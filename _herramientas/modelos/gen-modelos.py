# Genera los 10 modelos descargables: página HTML + Word (.docx) + PDF + Excel (.xlsx) cuando toca.
# Necesita python-docx y openpyxl (pip install python-docx openpyxl) y Google Chrome para los PDF.
# Si faltan, solo regenera las páginas HTML (los archivos ya publicados no se tocan).
# Uso: python3 _herramientas/modelos/gen-modelos.py
import html as H, json, os, re, subprocess, sys, tempfile
AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, '..'))
sys.path.insert(0, AQUI)
from plantilla import pagina, REVISADO, AD, RAIZ
from modelos import MODELOS, CASOS
try:
    import docx, openpyxl
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    ARCHIVOS = True                                   # sí podemos crear Word y Excel
except ImportError:
    ARCHIVOS = False
    print('Aviso: sin python-docx/openpyxl solo se regeneran las páginas')
CHROME = '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'
DESC = os.path.join(RAIZ, 'descargas')
os.makedirs(DESC, exist_ok=True)
HUECO = re.compile(r'(\[[^\]]+\])')                    # lo que va entre corchetes se rellena
REV = REVISADO.format(fecha='19 de octubre de 2026')
NOMBRE_FORMATO = {'docx': ('📝', 'Word'), 'pdf': ('📕', 'PDF'), 'xlsx': ('📊', 'Excel')}

# ================================================================ HTML del documento (página y PDF)
def con_huecos(t):
    return HUECO.sub(r'<span class="hueco">\1</span>', H.escape(t))

def doc_html(m, filas_registro=None):
    out = []
    for tipo, val in m['doc']:
        if tipo == 'titulo': out.append(f'<p class="doc-titulo">{con_huecos(val)}</p>')
        elif tipo == 'p': out.append(f'<p>{con_huecos(val)}</p>')
        elif tipo == 'datos': out.append('<p class="doc-datos">' + '<br>'.join(con_huecos(l) for l in val) + '</p>')
        elif tipo == 'lista': out.append('<ul>' + ''.join(f'<li>{con_huecos(l)}</li>' for l in val) + '</ul>')
        elif tipo == 'firma': out.append('<p class="doc-firma">' + '<br>'.join(con_huecos(l) or '&nbsp;' for l in val) + '</p>')
        elif tipo == 'tabla':
            filas = val
            if filas_registro:                         # en el PDF del registro: los 31 días del mes
                filas = [val[0]] + [[str(d)] + [''] * (len(val[0]) - 1) for d in range(1, 32)]
            cab = ''.join(f'<th>{H.escape(c)}</th>' for c in filas[0])
            cuerpo = ''.join('<tr>' + ''.join(f'<td>{con_huecos(c)}</td>' for c in f) + '</tr>' for f in filas[1:])
            out.append(f'<table class="doc-tabla"><thead><tr>{cab}</tr></thead><tbody>{cuerpo}</tbody></table>')
    return '\n'.join(out)

# ================================================================ Word
def crear_docx(m, ruta):
    d = docx.Document()
    estilo = d.styles['Normal']; estilo.font.name = 'Calibri'; estilo.font.size = Pt(11)
    for s in d.sections: s.left_margin = s.right_margin = Cm(2.5)
    def parrafo(texto, negrita=False, centro=False, junto=False):
        p = d.add_paragraph()
        if junto: p.paragraph_format.space_after = Pt(0)
        if centro: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for trozo in HUECO.split(texto):                # los huecos [..] van resaltados en amarillo
            if not trozo: continue
            r = p.add_run(trozo); r.bold = negrita
            if HUECO.fullmatch(trozo): r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p
    for tipo, val in m['doc']:
        if tipo == 'titulo': parrafo(val, negrita=True, centro=val.isupper())
        elif tipo == 'p': parrafo(val).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif tipo == 'datos':
            for i, l in enumerate(val): parrafo(l, negrita=(i == 0 and l.isupper()), junto=i < len(val) - 1)
        elif tipo == 'lista':
            for l in val:
                p = d.add_paragraph(style='List Bullet')
                for trozo in HUECO.split(l):
                    if trozo:
                        r = p.add_run(trozo)
                        if HUECO.fullmatch(trozo): r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif tipo == 'firma':
            for l in val: parrafo(l, junto=True)
            d.add_paragraph()
        elif tipo == 'tabla':
            filas = val if m['slug'] != 'registro-jornada' else [val[0]] + [[str(x)] + [''] * (len(val[0]) - 1) for x in range(1, 32)]
            t = d.add_table(rows=len(filas), cols=len(filas[0])); t.style = 'Table Grid'
            for i, f in enumerate(filas):
                for j, c in enumerate(f):
                    t.cell(i, j).text = c
                    if i == 0: t.cell(i, j).paragraphs[0].runs[0].bold = True
            d.add_paragraph()
    d.core_properties.author = ''; d.core_properties.comments = ''; d.core_properties.title = m['corto']   # sin datos ocultos raros al abrirlo
    d.save(ruta)

# ================================================================ Excel
if ARCHIVOS:                                          # estilos del Excel: solo si está openpyxl
    FINO = Side(style='thin', color='D1D5DB'); BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
    AMARILLO = PatternFill('solid', fgColor='FEF9C3'); CABEZA = PatternFill('solid', fgColor='EEF2FF')

def excel_factura(ruta):
    wb = openpyxl.Workbook(); wb.properties.creator = ''; ws = wb.active; ws.title = 'Factura'
    for col, ancho in zip('ABCD', (46, 12, 14, 16)): ws.column_dimensions[col].width = ancho
    ws['A1'] = 'FACTURA'; ws['A1'].font = Font(bold=True, size=18)
    datos = [('Número de factura', '2026-001'), ('Fecha de emisión', '01/01/2026'), ('', ''),
             ('EMISOR', ''), ('Nombre o razón social', '[Tu nombre]'), ('NIF', '[Tu NIF]'), ('Dirección', '[Tu dirección]'), ('', ''),
             ('CLIENTE', ''), ('Nombre o razón social', '[Cliente]'), ('NIF', '[NIF del cliente]'), ('Dirección', '[Dirección del cliente]')]
    for i, (k, v) in enumerate(datos, start=3):
        ws[f'A{i}'] = k; ws[f'B{i}'] = v
        if k.isupper(): ws[f'A{i}'].font = Font(bold=True)
        if v: ws[f'B{i}'].fill = AMARILLO
    f0 = 17                                            # fila de la cabecera de conceptos
    for j, c in enumerate(['Concepto', 'Cantidad', 'Precio (€)', 'Importe (€)']):
        cel = ws.cell(row=f0, column=j + 1, value=c); cel.font = Font(bold=True); cel.fill = CABEZA; cel.border = BORDE
    for r in range(f0 + 1, f0 + 9):                     # 8 líneas de conceptos
        for j in range(4): ws.cell(row=r, column=j + 1).border = BORDE
        for j in range(3): ws.cell(row=r, column=j + 1).fill = AMARILLO
        ws.cell(row=r, column=4, value=f'=IF(OR(B{r}="",C{r}=""),"",B{r}*C{r})').number_format = '#,##0.00 €'
        ws.cell(row=r, column=3).number_format = '#,##0.00 €'
    ws[f'A{f0 + 1}'] = 'Descripción del servicio'; ws[f'B{f0 + 1}'] = 1; ws[f'C{f0 + 1}'] = 500
    t = f0 + 10
    filas = [('Base imponible', f'=SUM(D{f0 + 1}:D{f0 + 8})'), ('Tipo de IVA', 0.21), ('IVA', f'=D{t}*D{t + 1}'),
             ('Retención de IRPF (0 % si el cliente es un particular)', 0.15), ('Retención', f'=-D{t}*D{t + 3}'), ('TOTAL A PAGAR', f'=D{t}+D{t + 2}+D{t + 4}')]
    for i, (k, v) in enumerate(filas):
        ws[f'A{t + i}'] = k; ws[f'D{t + i}'] = v
        ws[f'D{t + i}'].number_format = '0%' if isinstance(v, float) else '#,##0.00 €'
        if isinstance(v, float): ws[f'D{t + i}'].fill = AMARILLO
    ws[f'A{t + 5}'].font = ws[f'D{t + 5}'].font = Font(bold=True, size=13)
    ws[f'A{t + 7}'] = 'Forma de pago: transferencia a la cuenta [IBAN] antes del [fecha].'
    ws[f'A{t + 9}'] = 'Rellena solo las celdas amarillas: el resto se calcula solo.'
    ws[f'A{t + 9}'].font = Font(size=8, color='6B7280')
    wb.save(ruta)

def excel_registro(ruta):
    wb = openpyxl.Workbook(); wb.properties.creator = ''; ws = wb.active; ws.title = 'Registro'
    for col, ancho in zip('ABCDEFG', (8, 11, 13, 12, 11, 11, 30)): ws.column_dimensions[col].width = ancho
    ws['A1'] = 'REGISTRO DE JORNADA'; ws['A1'].font = Font(bold=True, size=16)
    for i, k in enumerate(['Empresa', 'Persona trabajadora', 'Mes y año'], start=3):
        ws[f'A{i}'] = k; ws[f'C{i}'] = f'[{k.lower()}]'; ws[f'C{i}'].fill = AMARILLO
    f0 = 7
    for j, c in enumerate(['Día', 'Entrada', 'Inicio pausa', 'Fin pausa', 'Salida', 'Horas', 'Observaciones']):
        cel = ws.cell(row=f0, column=j + 1, value=c); cel.font = Font(bold=True); cel.fill = CABEZA; cel.border = BORDE
        cel.alignment = Alignment(horizontal='center')
    for dia in range(1, 32):
        r = f0 + dia
        ws.cell(row=r, column=1, value=dia).alignment = Alignment(horizontal='center')
        for j in range(2, 6):
            c = ws.cell(row=r, column=j); c.number_format = 'hh:mm'; c.fill = AMARILLO
        # horas del día = (salida − entrada) − (fin pausa − inicio pausa); vacío si falta la entrada o la salida
        ws.cell(row=r, column=6, value=f'=IF(OR(B{r}="",E{r}=""),"",(E{r}-B{r})-(D{r}-C{r}))').number_format = '[h]:mm'
        for j in range(1, 8): ws.cell(row=r, column=j).border = BORDE
    from datetime import time
    ws.cell(row=f0 + 1, column=2, value=time(9)); ws.cell(row=f0 + 1, column=3, value=time(14))
    ws.cell(row=f0 + 1, column=4, value=time(15)); ws.cell(row=f0 + 1, column=5, value=time(18))
    tot = f0 + 33
    ws[f'E{tot}'] = 'Total del mes'; ws[f'E{tot}'].font = Font(bold=True)
    ws[f'F{tot}'] = f'=SUM(F{f0 + 1}:F{f0 + 31})'; ws[f'F{tot}'].number_format = '[h]:mm'; ws[f'F{tot}'].font = Font(bold=True)
    ws[f'A{tot + 2}'] = 'Firma de la persona trabajadora: ____________      Firma de la empresa: ____________'
    ws[f'A{tot + 4}'] = 'Escribe las horas como 9:00. Rellena solo las celdas amarillas.'
    ws[f'A{tot + 4}'].font = Font(size=8, color='6B7280')
    wb.save(ruta)

def excel_rectificativa(ruta):
    wb = openpyxl.Workbook(); wb.properties.creator = ''; ws = wb.active; ws.title = 'Rectificativa'
    for col, ancho in zip('ABCD', (40, 18, 18, 16)): ws.column_dimensions[col].width = ancho
    ws['A1'] = 'FACTURA RECTIFICATIVA'; ws['A1'].font = Font(bold=True, size=18)
    datos = [('Número (serie de rectificativas)', 'R-2026-001'), ('Fecha de emisión', '[fecha]'), ('Factura que se rectifica (n.º y fecha)', '[2026-015 de fecha …]'),
             ('Motivo', '[error en el importe / descuento / devolución…]'), ('', ''), ('EMISOR', ''), ('Nombre y NIF', '[Tu nombre · NIF]'),
             ('CLIENTE', ''), ('Nombre y NIF', '[Cliente · NIF]')]
    for i, (k, v) in enumerate(datos, start=3):
        ws[f'A{i}'] = k; ws[f'B{i}'] = v
        if k.isupper(): ws[f'A{i}'].font = Font(bold=True)
        if v: ws[f'B{i}'].fill = AMARILLO
    f0 = 14
    for j, c in enumerate(['Concepto', 'Importe anterior (€)', 'Importe correcto (€)', 'Diferencia (€)']):
        cel = ws.cell(row=f0, column=j + 1, value=c); cel.font = Font(bold=True); cel.fill = CABEZA; cel.border = BORDE
    for r in range(f0 + 1, f0 + 6):
        for j in range(4): ws.cell(row=r, column=j + 1).border = BORDE
        for j in range(3): ws.cell(row=r, column=j + 1).fill = AMARILLO
        for j in (2, 3): ws.cell(row=r, column=j).number_format = '#,##0.00 €'
        ws.cell(row=r, column=4, value=f'=IF(OR(B{r}="",C{r}=""),"",C{r}-B{r})').number_format = '#,##0.00 €'   # nuevo − anterior
    ws[f'A{f0 + 1}'] = 'Descripción'; ws[f'B{f0 + 1}'] = 1000; ws[f'C{f0 + 1}'] = 800
    t = f0 + 7
    filas = [('Base imponible (diferencia)', f'=SUM(D{f0 + 1}:D{f0 + 5})'), ('Tipo de IVA', 0.21), ('IVA', f'=D{t}*D{t + 1}'),
             ('Retención de IRPF (0 % si el cliente es un particular)', 0.15), ('Retención', f'=-D{t}*D{t + 3}'), ('TOTAL', f'=D{t}+D{t + 2}+D{t + 4}')]
    for i, (k, v) in enumerate(filas):
        ws[f'A{t + i}'] = k; ws[f'D{t + i}'] = v
        ws[f'D{t + i}'].number_format = '0%' if isinstance(v, float) else '#,##0.00 €'
        if isinstance(v, float): ws[f'D{t + i}'].fill = AMARILLO
    ws[f'A{t + 5}'].font = ws[f'D{t + 5}'].font = Font(bold=True, size=13)
    ws[f'A{t + 7}'] = 'Rellena solo las celdas amarillas: el resto se calcula solo. Si baja el importe, el total sale en negativo.'
    ws[f'A{t + 7}'].font = Font(size=8, color='6B7280')
    wb.save(ruta)

def excel_presupuesto(ruta):
    wb = openpyxl.Workbook(); wb.properties.creator = ''; ws = wb.active; ws.title = 'Presupuesto'
    for col, ancho in zip('ABCD', (46, 12, 14, 16)): ws.column_dimensions[col].width = ancho
    ws['A1'] = 'PRESUPUESTO'; ws['A1'].font = Font(bold=True, size=18)
    datos = [('Número', 'P-2026-001'), ('Fecha', '[fecha]'), ('Válido hasta', '[fecha]'), ('', ''), ('DE', ''), ('Nombre y NIF', '[Tu nombre · NIF]'),
             ('Contacto', '[teléfono · correo]'), ('', ''), ('PARA', ''), ('Cliente', '[Nombre del cliente]')]
    for i, (k, v) in enumerate(datos, start=3):
        ws[f'A{i}'] = k; ws[f'B{i}'] = v
        if k.isupper(): ws[f'A{i}'].font = Font(bold=True)
        if v: ws[f'B{i}'].fill = AMARILLO
    f0 = 15
    for j, c in enumerate(['Concepto', 'Cantidad', 'Precio (€)', 'Importe (€)']):
        cel = ws.cell(row=f0, column=j + 1, value=c); cel.font = Font(bold=True); cel.fill = CABEZA; cel.border = BORDE
    for r in range(f0 + 1, f0 + 9):
        for j in range(4): ws.cell(row=r, column=j + 1).border = BORDE
        for j in range(3): ws.cell(row=r, column=j + 1).fill = AMARILLO
        ws.cell(row=r, column=3).number_format = '#,##0.00 €'
        ws.cell(row=r, column=4, value=f'=IF(OR(B{r}="",C{r}=""),"",B{r}*C{r})').number_format = '#,##0.00 €'
    ws[f'A{f0 + 1}'] = 'Descripción del trabajo'; ws[f'B{f0 + 1}'] = 1; ws[f'C{f0 + 1}'] = 500
    t = f0 + 10
    filas = [('Base imponible', f'=SUM(D{f0 + 1}:D{f0 + 8})'), ('Tipo de IVA', 0.21), ('IVA', f'=D{t}*D{t + 1}'), ('TOTAL', f'=D{t}+D{t + 2}')]
    for i, (k, v) in enumerate(filas):
        ws[f'A{t + i}'] = k; ws[f'D{t + i}'] = v
        ws[f'D{t + i}'].number_format = '0%' if isinstance(v, float) else '#,##0.00 €'
        if isinstance(v, float): ws[f'D{t + i}'].fill = AMARILLO
    ws[f'A{t + 3}'].font = ws[f'D{t + 3}'].font = Font(bold=True, size=13)
    ws[f'A{t + 5}'] = 'Condiciones: [plazo de entrega] · [forma de pago]'; ws[f'A{t + 5}'].fill = AMARILLO
    ws[f'A{t + 7}'] = 'Aceptado por el cliente (firma y fecha): ______________________'
    wb.save(ruta)

def excel_nomina(ruta):
    wb = openpyxl.Workbook(); wb.properties.creator = ''; ws = wb.active; ws.title = 'Nómina'
    for col, ancho in zip('ABC', (44, 14, 16)): ws.column_dimensions[col].width = ancho
    ws['A1'] = 'RECIBO DE SALARIOS (orientativo)'; ws['A1'].font = Font(bold=True, size=16)
    for i, (k, v) in enumerate([('Empresa (nombre y CIF)', '[empresa]'), ('Persona trabajadora (nombre y DNI)', '[nombre]'), ('Periodo', '[mes y año]')], start=3):
        ws[f'A{i}'] = k; ws[f'B{i}'] = v; ws[f'B{i}'].fill = AMARILLO
    def fila(r, texto, valor, formato='#,##0.00 €', amarilla=False, negrita=False):
        ws[f'A{r}'] = texto; c = ws[f'C{r}']; c.value = valor; c.number_format = formato
        if amarilla: c.fill = AMARILLO
        if negrita: ws[f'A{r}'].font = c.font = Font(bold=True)
    ws['A7'] = 'DEVENGOS'; ws['A7'].font = Font(bold=True); ws['A7'].fill = CABEZA
    fila(8, 'Salario base', 1500, amarilla=True)
    fila(9, 'Complementos', 0, amarilla=True)
    fila(10, 'Horas extra', 0, amarilla=True)
    fila(11, 'Total devengado', '=SUM(C8:C10)', negrita=True)
    ws['A13'] = 'BASE DE COTIZACIÓN'; ws['A13'].font = Font(bold=True); ws['A13'].fill = CABEZA
    fila(14, '¿Pagas extra prorrateadas en la nómina? (escribe SI o NO)', 'NO', formato='@', amarilla=True)
    fila(15, 'Número de pagas extra al año', 2, formato='0', amarilla=True)
    fila(16, 'Parte de las pagas extra (si no van prorrateadas)', '=IF(UPPER(C14)="SI",0,(C8+C9)*C15/12)')
    fila(17, 'Base de contingencias comunes (sin horas extra)', '=C8+C9+C16', negrita=True)
    fila(18, 'Base de desempleo y formación (con horas extra)', '=C17+C10')
    ws['A19'] = 'DEDUCCIONES (tipos de 2026)'; ws['A19'].font = Font(bold=True); ws['A19'].fill = CABEZA
    # [concepto, tipo, base sobre la que se aplica]
    tipos = [('Contingencias comunes', 0.047, 'C17'), ('Desempleo (1,60 % si el contrato es temporal)', 0.0155, 'C18'),
             ('Formación profesional', 0.001, 'C18'), ('MEI', 0.0015, 'C17'), ('Horas extra (cotización propia)', 0.047, 'C10')]
    for i, (k, tp, b) in enumerate(tipos):
        r = 20 + i; ws[f'A{r}'] = k; ws[f'B{r}'] = tp; ws[f'B{r}'].number_format = '0.00%'
        if i == 1: ws[f'B{r}'].fill = AMARILLO
        ws[f'C{r}'] = f'={b}*B{r}'; ws[f'C{r}'].number_format = '#,##0.00 €'
    ws['A25'] = 'IRPF (pon tu porcentaje)'; ws['B25'] = 0.12; ws['B25'].number_format = '0.00%'; ws['B25'].fill = AMARILLO
    ws['C25'] = '=C11*B25'; ws['C25'].number_format = '#,##0.00 €'
    fila(26, 'Total deducciones', '=SUM(C20:C25)', negrita=True)
    fila(28, 'LÍQUIDO A PERCIBIR', '=C11-C26', negrita=True)
    ws['A28'].font = ws['C28'].font = Font(bold=True, size=14)
    ws['A30'] = 'Orientativa: no incluye casos especiales (bajas, embargos, atrasos, horas extra por fuerza mayor…).'
    ws['A30'].font = Font(size=8, color='6B7280')
    wb.save(ruta)

EXCEL = {'factura-autonomo': excel_factura, 'registro-jornada': excel_registro, 'factura-rectificativa': excel_rectificativa,
         'presupuesto-autonomo': excel_presupuesto, 'plantilla-nomina': excel_nomina}

# ================================================================ PDF (con Chrome, sin ventanas)
def crear_pdf(m, ruta):
    cuerpo = doc_html(m, filas_registro=(m['slug'] == 'registro-jornada'))
    pagina_html = f'''<!doctype html><html><head><meta charset="utf-8"><style>
    @page {{ size: A4; margin: 22mm 20mm; }}
    body {{ font-family: Calibri, "Helvetica Neue", Arial, sans-serif; font-size: 11pt; line-height: 1.5; color: #111; }}
    p {{ margin: 0 0 10pt; text-align: justify; }} .doc-titulo {{ font-weight: 700; }} .doc-datos {{ text-align: left; }}
    .hueco {{ background: #FEF08A; }} .doc-firma {{ margin-top: 14pt; text-align: left; }}
    table {{ width: 100%; border-collapse: collapse; margin: 6pt 0 12pt; }} th, td {{ border: 1px solid #999; padding: 3pt 5pt; font-size: 10pt; text-align: left; }}
    th {{ background: #EEF2FF; }} .aviso {{ font-size: 8pt; color: #6B7280; margin-top: 18pt; }}
    {'@page { margin: 12mm 16mm; } td, th { padding: 1pt 5pt; font-size: 9pt; line-height: 1.3; } body { font-size: 10pt; line-height: 1.35; } p { margin-bottom: 5pt; } .doc-firma { margin-top: 8pt; } .aviso { margin-top: 8pt; }' if m['slug'] == 'registro-jornada' else ''}
    </style><title>{H.escape(m['corto'])}</title></head><body>{cuerpo}</body></html>'''
    with tempfile.NamedTemporaryFile('w', suffix='.html', delete=False, encoding='utf-8') as t:
        t.write(pagina_html)
    subprocess.run([CHROME, '--headless=new', '--disable-gpu', '--no-pdf-header-footer', f'--print-to-pdf={ruta}', 'file://' + t.name],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
    os.unlink(t.name)

def peso(ruta):
    return f'{max(1, round(os.path.getsize(ruta) / 1024))} KB' if os.path.exists(ruta) else ''

# ================================================================ Páginas
gen = []
FUENTE_GEN = 'Modelo orientativo redactado por CalculaLoTuyo a partir del texto vigente del ' + '<a href="https://www.boe.es/buscar/act.php?id=BOE-A-2015-11430" target="_blank">Estatuto de los Trabajadores</a>. Revisa siempre tu convenio colectivo.'
for m in MODELOS:
    base = os.path.join(DESC, m['slug'])
    if ARCHIVOS:
        for f in m['formatos']:
            if f == 'docx': crear_docx(m, base + '.docx')
            if f == 'pdf': crear_pdf(m, base + '.pdf')
            if f == 'xlsx': EXCEL[m['slug']](base + '.xlsx')
    botones = ''.join(f'<a href="descargas/{m["slug"]}.{f}" class="boton-descarga" download><span aria-hidden="true">{NOMBRE_FORMATO[f][0]}</span><strong>Descargar en {NOMBRE_FORMATO[f][1]}</strong><small>.{f} · {peso(base + "." + f)}</small></a>' for f in m['formatos'])
    extra = ''
    if m.get('tabla_extra'):
        t = m['tabla_extra']
        extra = '<h2 class="puentes-titulo">📅 Días de permiso según el motivo</h2><div class="tabla-wrap"><table class="tabla-festivos"><thead><tr>' + ''.join(f'<th>{c}</th>' for c in t[0]) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join(f'<td>{c}</td>' for c in f) + '</tr>' for f in t[1:]) + '</tbody></table></div>'
    otros = '<div class="enlaces-sueldo">' + ''.join(f'<a href="modelo-{o["slug"]}">{o["icono"]} {o["corto"]}</a>' for o in MODELOS if o is not m) + '</div>'
    rel = '<div class="enlaces-sueldo">' + ''.join(f'<a href="{u}">{t}</a>' for u, t in m['rel']) + '</div>'
    pasos = ''.join(f'<div class="calc-step"><span class="step-num">{i}</span><div><strong>{t}</strong><p>{x}</p></div></div>' for i, (t, x) in enumerate(m['pasos'], 1))
    errores = ''.join(f'<li>{e}</li>' for e in m['errores'])
    extra_kw, lista_casos = CASOS.get(m['slug'], ('', []))
    casos = ('<h2>💡 Casos frecuentes</h2>' + ''.join(f'<h3>{t}</h3><p>{x}</p>' for t, x in lista_casos)) if lista_casos else ''
    cuerpo = f'''    <section class="hero-mini"><div class="container">{{{{MIGAS}}}}<div class="hero-badge">📄 Modelo gratis</div><h1>{m["h1"]}<br><span class="highlight">{m["sub"]}</span></h1><p>{m["intro"]}</p>
            {REV}</div></section>

    <section class="calc-section"><div class="container"><div class="card">
        <div class="descargas">{botones}</div>
        <p class="nota-corta descargas-nota">Gratis y sin registro. Lo que está en <span class="hueco">amarillo</span> es lo que tienes que rellenar.</p>
        {extra}
        <div class="doc-cabecera"><h2 class="puentes-titulo">👀 Así es el modelo</h2><button type="button" class="btn-secondary" id="copiar">📋 Copiar el texto</button></div>
        <div class="documento" id="documento">{doc_html(m)}</div>
    </div></div></section>
''' + AD + f'''
    <section class="section"><div class="container"><article class="seo-section">
<h2>✍️ Cómo rellenarlo</h2>
<div class="calc-steps">{pasos}</div>
{casos}
<h2>⚠️ Errores frecuentes</h2>
<ul>{errores}</ul>
<h2>⚖️ Qué dice la ley</h2>
<p>{m["base"]} Es un modelo orientativo: tu convenio colectivo puede mejorar lo que dice la ley.</p>
<h2>🔗 Te puede servir</h2>
{rel}
<h2>📄 Otros modelos gratis</h2>
{otros}
</article></div></section>
<script>
// Botón «Copiar el texto»: copia el modelo tal cual se ve, para pegarlo en un correo o en Word
document.getElementById('copiar').addEventListener('click',async function(){{
    const texto=document.getElementById('documento').innerText;   // el texto visible del modelo
    try{{await navigator.clipboard.writeText(texto);this.textContent='✅ Copiado'}}   // lo copia al portapapeles
    catch(e){{this.textContent='Selecciona y copia a mano'}}         // por si el navegador no lo permite
    setTimeout(()=>{{this.textContent='📋 Copiar el texto'}},2500);  // vuelve al texto original
}});
</script>
'''
    gen.append(pagina(f'modelo-{m["slug"]}', m['titulo'], m['desc'], m['kw'] + extra_kw, [('Modelos', 'modelos'), (m['corto'], None)], cuerpo, m['faqs'], FUENTE_GEN, ''))

# ================================================================ Página «Modelos»
tarjetas = ''.join(f'<a href="modelo-{m["slug"]}" class="modelo-card"><span class="modelo-icono" aria-hidden="true">{m["icono"]}</span><strong>{m["corto"]}</strong><span>{m["intro"]}</span><small>{" · ".join(NOMBRE_FORMATO[f][1] for f in m["formatos"])}</small></a>' for m in MODELOS)
cuerpo = f'''    <section class="hero-mini"><div class="container">{{{{MIGAS}}}}<div class="hero-badge">📄 Modelos gratis</div><h1>Modelos y plantillas laborales<br><span class="highlight">gratis en Word, PDF y Excel</span></h1><p>Cartas y plantillas listas para rellenar: baja voluntaria, reducción de jornada, vacaciones, facturas… Revisadas con el Estatuto de los Trabajadores.</p>
            {REV}</div></section>

    <section class="calc-section"><div class="container"><div class="card">
        <div class="modelos-grid">{tarjetas}</div>
        <p class="nota-corta">Todos gratis y sin registro. Son orientativos: revisa tu convenio colectivo, que puede mejorar lo que dice la ley.</p>
    </div></div></section>
''' + AD
faqs = [('¿Son gratis los modelos?', 'Sí, <strong>todos son gratis y sin registro</strong>. Puedes descargarlos en Word, PDF o Excel, o copiar el texto directamente desde la página.'),
        ('¿Sirven para cualquier empresa?', 'Están redactados con el Estatuto de los Trabajadores, que se aplica a todas las personas con contrato laboral en España. Tu convenio colectivo puede mejorar plazos o días, así que conviene revisarlo.'),
        ('¿Cómo entrego el escrito a la empresa?', 'En mano con dos copias (y que te firmen la tuya con la fecha), por correo electrónico con acuse de recibo o por burofax si quieres una prueba más fuerte.')]
gen.append(pagina('modelos', 'Modelos y plantillas laborales gratis (Word, PDF y Excel)', 'Modelos gratis de carta de baja voluntaria, reducción de jornada, vacaciones, excedencia, reclamación de finiquito, factura de autónomo y registro de jornada. Word, PDF y Excel.', 'modelos laborales gratis, plantillas laborales word, carta baja voluntaria, modelo reducción de jornada, plantilla factura autónomo', [('Modelos', None)], cuerpo, faqs, FUENTE_GEN, ''))
json.dump(gen, open(f'{AQUI}/generadas.json', 'w'))
print(len(gen), 'páginas ·', 'archivos creados' if ARCHIVOS else 'sin archivos')
